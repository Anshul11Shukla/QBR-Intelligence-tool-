
from io import BytesIO

import pytest
from pptx import Presentation

from app.services.pptx_service import generate_qbr_pptx


@pytest.fixture
def sample_context():
    return {
        "executive_summary": "Quarterly service performance reviewed.",
        "kpi_insights": [
            "SLA Compliance decreased quarter-on-quarter.",
            "CSAT improved quarter-on-quarter.",
        ],
        "kpi_evidence": [
            {
                "metric": "SLA Compliance",
                "previous_value": 98,
                "current_value": 92,
                "target": 95,
                "sources": {
                    "previous": "previous_quarter.xlsx / Sheet1",
                    "current": "current_quarter.xlsx / Sheet1",
                },
            }
        ],
        "prioritized_risks": [
            {
                "metric": "SLA Compliance",
                "severity": "high",
                "reason": "KPI performance declined",
            }
        ],
        "commitments": [
            {
                "title": "Update incident response playbook",
                "owner": "Operations Team",
                "status": "in_progress",
                "effective_status": "overdue",
            }
        ],
        "recommendations": [
            "Review SLA breach incidents and escalation delays."
        ],
        "questions": [
            {
                "customer_question": "Why did SLA Compliance decline?",
                "suggested_answer": "The metric decreased from 98% to 92%.",
                "preparation_actions": [
                    "Review incident records."
                ],
                "answer_readiness": "partially_ready",
            }
        ],
    }


def get_slide_text(slide):
    return "\n".join(
        shape.text
        for shape in slide.shapes
        if shape.has_text_frame
    )


def test_pptx_generates_valid_file(sample_context):
    result = generate_qbr_pptx(
        sample_context,
        "ABC Technologies",
        "Q3 2026",
    )

    assert isinstance(result, bytes)
    assert len(result) > 1000

    presentation = Presentation(BytesIO(result))
    assert len(presentation.slides) >= 8


def test_pptx_contains_customer_name(sample_context):
    result = generate_qbr_pptx(
        sample_context,
        "ABC Technologies",
        "Q3 2026",
    )

    presentation = Presentation(BytesIO(result))
    first_slide_text = get_slide_text(presentation.slides[0])

    assert "ABC Technologies" in first_slide_text
    assert "Q3 2026" in first_slide_text


def test_pptx_contains_kpi_values(sample_context):
    result = generate_qbr_pptx(sample_context)
    presentation = Presentation(BytesIO(result))

    table_values = []

    for slide in presentation.slides:
        for shape in slide.shapes:
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        table_values.append(cell.text)

    assert "SLA Compliance" in table_values
    assert "98" in table_values
    assert "92" in table_values
    assert "95" in table_values


def test_pptx_contains_customer_questions(sample_context):
    result = generate_qbr_pptx(sample_context)
    presentation = Presentation(BytesIO(result))

    all_text = "\n".join(
        get_slide_text(slide)
        for slide in presentation.slides
    )

    assert "Why did SLA Compliance decline?" in all_text
    assert "The metric decreased from 98% to 92%." in all_text


def test_pptx_rejects_invalid_context():
    with pytest.raises(ValueError):
        generate_qbr_pptx(None)
