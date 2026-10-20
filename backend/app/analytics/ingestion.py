
from io import BytesIO

from openpyxl import load_workbook

from app.analytics.normalization import (
    metric_key,
    normalize_direction,
    normalize_metric,
    normalize_unit,
    parse_number,
)
from app.schemas.kpi import KPIRecord


REQUIRED_COLUMNS = {
    "metric",
    "value",
    "unit",
    "direction",
    "target",
}

MAX_ROWS = 5000


def parse_workbook(content: bytes) -> list[KPIRecord]:
    try:
        workbook = load_workbook(
            BytesIO(content),
            read_only=True,
            data_only=True,
        )
    except Exception as exc:
        raise ValueError(
            "Invalid or unreadable XLSX workbook"
        ) from exc

    records = []
    seen = set()

    try:
        for sheet in workbook.worksheets:
            rows = sheet.iter_rows(values_only=True)
            header = next(rows, None)

            if header is None:
                continue

            columns = [
                str(value or "").strip().lower()
                for value in header
            ]

            if len(columns) != len(set(columns)):
                raise ValueError(
                    f"{sheet.title}: duplicate column headings"
                )

            if not REQUIRED_COLUMNS.issubset(columns):
                missing = REQUIRED_COLUMNS - set(columns)
                raise ValueError(
                    f"{sheet.title}: missing columns: "
                    + ", ".join(sorted(missing))
                )

            indexes = {
                name: columns.index(name)
                for name in REQUIRED_COLUMNS
            }

            for row_number, row in enumerate(rows, start=2):
                if row_number > MAX_ROWS + 1:
                    raise ValueError(
                        f"{sheet.title}: too many rows"
                    )

                if not any(
                    value is not None and str(value).strip()
                    for value in row
                ):
                    continue

                def cell(name):
                    index = indexes[name]
                    return row[index] if index < len(row) else None

                try:
                    metric = normalize_metric(cell("metric"))
                    key = metric_key(metric)

                    if key in seen:
                        raise ValueError(
                            f"Duplicate metric: {metric}"
                        )

                    record = KPIRecord(
                        metric=metric,
                        value=parse_number(cell("value")),
                        unit=normalize_unit(cell("unit")),
                        direction=normalize_direction(
                            cell("direction")
                        ),
                        target=parse_number(cell("target")),
                        source_sheet=sheet.title,
                        source_row=row_number,
                    )

                    seen.add(key)
                    records.append(record)

                except ValueError as exc:
                    raise ValueError(
                        f"{sheet.title}, row {row_number}: {exc}"
                    ) from exc

        if not records:
            raise ValueError("Workbook contains no KPI records")

        return records
    finally:
        workbook.close()
