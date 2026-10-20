
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

from docx import Document
from pypdf import PdfReader


@dataclass
class ParsedSection:
    location: str
    text: str


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}


def parse_document(
    filename: str,
    content: bytes,
) -> list[ParsedSection]:
    extension = Path(filename).suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            "Supported formats: PDF, DOCX, TXT, MD"
        )

    sections = []

    if extension == ".pdf":
        try:
            reader = PdfReader(BytesIO(content))

            if reader.is_encrypted:
                raise ValueError(
                    "Password-protected PDFs are not supported"
                )

            if len(reader.pages) > 100:
                raise ValueError("PDF exceeds 100 pages")

            for index, page in enumerate(reader.pages, start=1):
                text = (page.extract_text() or "").strip()

                if text:
                    sections.append(
                        ParsedSection(
                            location=f"Page {index}",
                            text=text,
                        )
                    )

        except ValueError:
            raise
        except Exception as exc:
            raise ValueError(
                "Could not parse PDF document"
            ) from exc

    elif extension == ".docx":
        try:
            document = Document(BytesIO(content))

            for index, paragraph in enumerate(
                document.paragraphs, start=1
            ):
                text = paragraph.text.strip()

                if text:
                    sections.append(
                        ParsedSection(
                            location=f"Paragraph {index}",
                            text=text,
                        )
                    )

            for table_index, table in enumerate(
                document.tables, start=1
            ):
                for row_index, row in enumerate(
                    table.rows, start=1
                ):
                    text = " | ".join(
                        cell.text.strip()
                        for cell in row.cells
                    )

                    if text.strip(" |"):
                        sections.append(
                            ParsedSection(
                                location=(
                                    f"Table {table_index}, "
                                    f"row {row_index}"
                                ),
                                text=text,
                            )
                        )

        except Exception as exc:
            raise ValueError(
                "Could not parse DOCX document"
            ) from exc

    else:
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ValueError(
                "Text files must use UTF-8 encoding"
            ) from exc

        for index, line in enumerate(
            text.splitlines(), start=1
        ):
            if line.strip():
                sections.append(
                    ParsedSection(
                        location=f"Line {index}",
                        text=line.strip(),
                    )
                )

    if not sections:
        raise ValueError(
            "No extractable text found in document"
        )

    return sections
