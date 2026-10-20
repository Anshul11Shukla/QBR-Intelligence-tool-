
from io import BytesIO
from typing import Any

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor


NAVY = RGBColor(20, 36, 65)
BLUE = RGBColor(37, 99, 235)
WHITE = RGBColor(255, 255, 255)
LIGHT = RGBColor(241, 245, 249)
DARK = RGBColor(51, 65, 85)
RED = RGBColor(185, 28, 28)
AMBER = RGBColor(180, 83, 9)

WIDTH = 13.33
HEIGHT = 7.5


def as_dict(value: Any) -> dict:
    if isinstance(value, dict):
        return value
    if hasattr(value, "model_dump"):
        return value.model_dump()
    return {}


def display(value: Any, default="Not available") -> str:
    if value is None or value == "":
        return default
    return str(value)


def add_text(slide, text, x, y, w, h, size=16,
             color=DARK, bold=False):
    shape = slide.shapes.add_textbox(
        Inches(x), Inches(y), Inches(w), Inches(h)
    )
    frame = shape.text_frame
    frame.word_wrap = True
    frame.margin_left = Inches(0.08)
    frame.margin_top = Inches(0.05)

    paragraph = frame.paragraphs[0]
    paragraph.text = display(text, "")
    paragraph.font.name = "Aptos"
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = color
    return shape


def new_slide(presentation, title):
    slide = presentation.slides.add_slide(
        presentation.slide_layouts[6]
    )

    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = WHITE

    bar = slide.shapes.add_shape(
        1, 0, 0, presentation.slide_width, Inches(0.13)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = BLUE
    bar.line.fill.background()

    add_text(slide, title, 0.65, 0.35, 12, 0.65,
             size=27, color=NAVY, bold=True)

    add_text(
        slide,
        "QBR Intelligence Studio  |  Evidence-grounded review",
        0.7, 7.12, 11.5, 0.22,
        size=10, color=DARK
    )

    return slide


def add_items(presentation, title, items, per_slide=5):
    items = [display(item) for item in items if item is not None]
    if not items:
        items = ["No verified information available."]

    for start in range(0, len(items), per_slide):
        page = start // per_slide + 1
        heading = title if page == 1 else f"{title} (continued)"
        slide = new_slide(presentation, heading)

        selected = items[start:start + per_slide]
        step = min(1.08, 5.7 / max(len(selected), 1))

        for index, item in enumerate(selected):
            y = 1.35 + index * step
            add_text(
                slide, f"{start + index + 1}. {item}",
                0.85, y, 11.6, step - 0.08,
                size=15
            )


def add_kpi_table(presentation, evidence):
    rows = [as_dict(item) for item in evidence]
    if not rows:
        add_items(
            presentation, "KPI Performance",
            ["No verified KPI comparison data available."]
        )
        return

    for start in range(0, len(rows), 7):
        slide = new_slide(
            presentation,
            "KPI Performance — Quarter Comparison"
        )
        selected = rows[start:start + 7]

        table = slide.shapes.add_table(
            len(selected) + 1, 4,
            Inches(0.65), Inches(1.45),
            Inches(12), Inches(0.65 * (len(selected) + 1))
        ).table

        table.columns[0].width = Inches(4.1)
        table.columns[1].width = Inches(2.65)
        table.columns[2].width = Inches(2.65)
        table.columns[3].width = Inches(2.6)

        headers = [
            "KPI", "Previous Quarter",
            "Current Quarter", "Target"
        ]

        for column, header in enumerate(headers):
            table.cell(0, column).text = header

        for row_index, item in enumerate(selected, 1):
            values = [
                display(item.get("metric")),
                display(item.get("previous_value")),
                display(item.get("current_value")),
                display(item.get("target")),
            ]

            for column, value in enumerate(values):
                table.cell(row_index, column).text = value

        for row_index in range(len(selected) + 1):
            for column in range(4):
                cell = table.cell(row_index, column)
                cell.fill.solid()
                cell.fill.fore_color.rgb = (
                    NAVY if row_index == 0
                    else LIGHT if row_index % 2 == 0
                    else WHITE
                )

                for paragraph in cell.text_frame.paragraphs:
                    paragraph.font.name = "Aptos"
                    paragraph.font.size = Pt(15)
                    paragraph.font.bold = row_index == 0
                    paragraph.font.color.rgb = (
                        WHITE if row_index == 0 else DARK
                    )

        add_text(
            slide,
            "Values are taken from uploaded KPI evidence. "
            "Target status depends on metric direction.",
            0.75, 6.6, 11.8, 0.4, size=11
        )


def add_qa_slides(presentation, questions):
    if not questions:
        add_items(
            presentation, "Customer Q&A Preparation",
            ["No customer questions available."]
        )
        return

    for index, raw_question in enumerate(questions, 1):
        question = as_dict(raw_question)
        slide = new_slide(
            presentation,
            f"Customer Q&A Preparation — Question {index}"
        )

        add_text(
            slide,
            display(question.get("customer_question")),
            0.85, 1.35, 11.6, 1.05,
            size=21, color=NAVY, bold=True
        )

        add_text(
            slide, "Suggested Answer",
            0.85, 2.55, 11.5, 0.4,
            size=16, color=BLUE, bold=True
        )

        add_text(
            slide,
            display(question.get("suggested_answer")),
            0.85, 3.0, 11.6, 2.15,
            size=15
        )

        actions = question.get("preparation_actions", [])
        if isinstance(actions, list) and actions:
            add_text(
                slide, "Preparation Action",
                0.85, 5.35, 11.5, 0.35,
                size=15, color=BLUE, bold=True
            )
            add_text(
                slide, display(actions[0]),
                0.85, 5.8, 11.6, 0.65,
                size=13
            )

        add_text(
            slide,
            "Answer readiness: "
            + display(question.get("answer_readiness")),
            0.85, 6.65, 11.5, 0.3,
            size=11
        )


def generate_qbr_pptx(
    context: dict,
    customer_name: str = "Customer",
    period: str = "Quarterly Business Review",
) -> bytes:

    if not isinstance(context, dict):
        raise ValueError("QBR context must be a dictionary")

    presentation = Presentation()
    presentation.slide_width = Inches(WIDTH)
    presentation.slide_height = Inches(HEIGHT)

    # 1. Cover
    slide = new_slide(presentation, "Quarterly Business Review")
    add_text(
        slide, customer_name,
        0.9, 2.3, 11.5, 1.1,
        size=37, color=BLUE, bold=True
    )
    add_text(
        slide, period,
        0.9, 3.55, 11.5, 0.6,
        size=23, color=NAVY
    )

    # 2. Executive Summary
    summary = context.get("executive_summary")
    add_items(
        presentation,
        "Executive Summary",
        [summary] if isinstance(summary, str) else (
            summary if isinstance(summary, list) else []
        ),
        per_slide=3
    )

    # 3. KPI Comparison
    add_kpi_table(
        presentation,
        context.get("kpi_evidence") or []
    )

    # 4. KPI Insights
    add_items(
        presentation,
        "Quarter-on-Quarter Insights",
        context.get("kpi_insights") or []
    )

    # 5. Prioritized Risks
    risks = []
    for raw_risk in context.get("prioritized_risks") or []:
        risk = as_dict(raw_risk)
        severity = display(risk.get("severity"), "unspecified").upper()
        metric = display(risk.get("metric"))
        reason = display(risk.get("reason"))

        risks.append(f"[{severity}] {metric}: {reason}")

    add_items(
        presentation, "Prioritized Risks",
        risks, per_slide=5
    )

    # 6. Commitments
    commitments = []
    for raw_item in context.get("commitments") or []:
        item = as_dict(raw_item)
        commitments.append(
            f"{display(item.get('title'))} | "
            f"Owner: {display(item.get('owner'))} | "
            f"Status: {display(item.get('effective_status') or item.get('status'))}"
        )

    add_items(
        presentation,
        "Historical Commitments",
        commitments
    )

    # 7. Recommendations
    recommendations = context.get("recommendations") or []
    add_items(
        presentation,
        "Recommendations and Next Steps",
        recommendations
    )

    # 8. Customer Q&A
    add_qa_slides(
        presentation,
        context.get("questions") or []
    )

    # 9. Evidence References
    references = []
    for raw_item in context.get("kpi_evidence") or []:
        item = as_dict(raw_item)
        sources = as_dict(item.get("sources"))
        metric = display(item.get("metric"))

        references.append(
            f"{metric}\n"
            f"Previous: {display(sources.get('previous'))}\n"
            f"Current: {display(sources.get('current'))}"
        )

    add_items(
        presentation,
        "KPI Evidence and References",
        references, per_slide=3
    )

    output = BytesIO()
    presentation.save(output)
    return output.getvalue()
